from io import BytesIO
from typing import Mapping

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


FINANCE_ROW_STYLES = {
    "total": {
        "fill": "B71C1C",
        "font_color": "FFFFFF",
        "bold": True,
    },
    "overdue": {
        "fill": "FDE2E2",
        "font_color": "7A1F1F",
        "bold": False,
    },
    "fact": {
        "fill": "EAF7EA",
        "font_color": "000000",
        "bold": False,
    },
    "quarter_1": {
        "fill": "EEF5FF",
        "font_color": "000000",
        "bold": False,
    },
    "quarter_2": {
        "fill": "EEFAF1",
        "font_color": "000000",
        "bold": False,
    },
    "quarter_3": {
        "fill": "FFF8E6",
        "font_color": "000000",
        "bold": False,
    },
    "quarter_4": {
        "fill": "F7EEFC",
        "font_color": "000000",
        "bold": False,
    },
    "default": {
        "fill": None,
        "font_color": "000000",
        "bold": False,
    },
}

HEADER_FILL = "1F4E78"
HEADER_FONT_COLOR = "FFFFFF"
BORDER_COLOR = "D9E2F3"
OVERDUE_STATUSES = {
    "⚠️ Плановая дата прошла, факта нет",
    "🔴 Плановая дата прошла, факта нет",
}


def get_finance_row_style(row: Mapping) -> dict:
    """Возвращает единый semantic style для Streamlit и Excel."""
    if str(row.get("Год", "")).endswith("ИТОГО"):
        return FINANCE_ROW_STYLES["total"]

    status = row.get("Статус", "")

    if status in OVERDUE_STATUSES:
        return FINANCE_ROW_STYLES["overdue"]

    if status == "✅ Факт":
        return FINANCE_ROW_STYLES["fact"]

    try:
        quarter = int(row.get("Квартал", 0))
    except (TypeError, ValueError):
        quarter = 0

    return FINANCE_ROW_STYLES.get(
        f"quarter_{quarter}",
        FINANCE_ROW_STYLES["default"],
    )


def _excel_value(value):
    if pd.isna(value):
        return None

    if hasattr(value, "item"):
        return value.item()

    return value


def export_finance_table(df: pd.DataFrame, unit: str = "шт") -> bytes:
    """Создаёт форматированный XLSX финансового плана поставок в памяти."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "План поставок"
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False

    headers = df.columns.tolist()
    worksheet.append(headers)

    for row in df.itertuples(index=False, name=None):
        worksheet.append([_excel_value(value) for value in row])

    thin_side = Side(style="thin", color=f"FF{BORDER_COLOR}")
    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )
    header_fill = PatternFill("solid", fgColor=f"FF{HEADER_FILL}")

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color=f"FF{HEADER_FONT_COLOR}", bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = thin_border

    worksheet.row_dimensions[1].height = 30

    base_columns = {
        "Год": 14,
        "Квартал": 12,
        "Месяц": 10,
        "№ недели": 12,
        "№ конт": 14,
        "Статус": 38,
    }
    quantity_columns = [
        column
        for column in headers
        if column not in base_columns
    ]
    quantity_format = "#,##0" if unit == "шт" else "#,##0.00"

    for column_index, column_name in enumerate(headers, start=1):
        column_letter = get_column_letter(column_index)
        worksheet.column_dimensions[column_letter].width = base_columns.get(
            column_name,
            min(max(len(str(column_name)) + 2, 10), 18),
        )

    for row_index, (_, row) in enumerate(df.iterrows(), start=2):
        style = get_finance_row_style(row)
        fill = (
            PatternFill("solid", fgColor=f"FF{style['fill']}")
            if style["fill"]
            else PatternFill(fill_type=None)
        )

        for column_index, column_name in enumerate(headers, start=1):
            cell = worksheet.cell(row=row_index, column=column_index)
            cell.fill = fill
            cell.font = Font(
                color=f"FF{style['font_color']}",
                bold=style["bold"],
            )
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal=(
                    "right"
                    if column_name in quantity_columns
                    else "left" if column_name == "Статус" else "center"
                ),
                vertical="center",
                wrap_text=column_name == "Статус",
            )

            if column_name in quantity_columns:
                cell.number_format = quantity_format

        worksheet.row_dimensions[row_index].height = 20

    if headers:
        last_column = get_column_letter(len(headers))
        worksheet.auto_filter.ref = f"A1:{last_column}{worksheet.max_row}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
