from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RISK_FILL_COLORS = {
    "Свежий 🟢": "2ECC71",
    "Хороший 🟢": "6FCF97",
    "Норма 🟡": "F1C40F",
    "Внимание 🟡": "F4D03F",
    "Пограничный 🟠": "F39C12",
    "Горящий 🔴": "E74C3C",
    "Критично 🔴": "C0392B",
    "Списание ⚫": "2C2C2C",
    "Ошибка": "9B59B6",
}

QUANTITY_COLUMNS = {
    "Остаток",
    "Зарезервировано",
    "Свободный остаток",
}
DATE_COLUMNS = {"Срок годности"}
PERCENT_COLUMNS = {"ОСГ %", "min ОСГ %", "Доля резерва, %"}


def _prepare_export_data(df: pd.DataFrame) -> pd.DataFrame:
    export_df = df.copy()
    for column in DATE_COLUMNS.intersection(export_df.columns):
        export_df[column] = pd.to_datetime(
            export_df[column],
            format="%d.%m.%Y",
            errors="coerce",
        ).dt.date
    return export_df


def _risk_values(
        df: pd.DataFrame,
        risk_values: pd.Series | None,
) -> list[str | None]:
    if risk_values is not None:
        aligned = risk_values.reindex(df.index)
        return aligned.astype("object").where(aligned.notna(), None).tolist()
    if "Риск" in df.columns:
        return df["Риск"].astype("object").where(df["Риск"].notna(), None).tolist()
    return [None] * len(df)


def export_osg_table(
        df: pd.DataFrame,
        sheet_name: str,
        risk_values: pd.Series | None = None,
) -> bytes:
    """Creates a formatted XLSX representation of an OSG user table."""
    export_df = _prepare_export_data(df)
    row_risks = _risk_values(df, risk_values)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]
    worksheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for column_index, column_name in enumerate(export_df.columns, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=column_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, row in enumerate(export_df.itertuples(index=False, name=None), start=2):
        risk = row_risks[row_index - 2]
        risk_fill = RISK_FILL_COLORS.get(risk)
        for column_index, value in enumerate(row, start=1):
            if pd.isna(value):
                value = None
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            column_name = export_df.columns[column_index - 1]

            if column_name in QUANTITY_COLUMNS:
                cell.number_format = "#,##0"
            elif column_name in DATE_COLUMNS:
                cell.number_format = "DD.MM.YYYY"
            elif column_name in PERCENT_COLUMNS:
                cell.number_format = '0.##"%"'

            if risk_fill and (column_name in PERCENT_COLUMNS or column_name == "Риск"):
                cell.fill = PatternFill("solid", fgColor=risk_fill)
                if risk in {"Горящий 🔴", "Критично 🔴", "Списание ⚫"}:
                    cell.font = Font(color="FFFFFF")

    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 24

    for column_index, column_name in enumerate(export_df.columns, start=1):
        values = [str(column_name)]
        values.extend(
            "" if value is None or pd.isna(value) else str(value)
            for value in export_df.iloc[:, column_index - 1]
        )
        width = min(max(max(map(len, values)) + 2, 11), 45)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
