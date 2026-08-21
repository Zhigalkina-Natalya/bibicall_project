import unittest
from io import BytesIO

from openpyxl import load_workbook
import pandas as pd

from src.arrivals.exporter import (
    BORDER_COLOR,
    FINANCE_ROW_STYLES,
    HEADER_FILL,
    export_finance_table,
    get_finance_row_style,
)


class ArrivalsFinanceExporterTests(unittest.TestCase):
    def setUp(self):
        self.df = pd.DataFrame({
            "Год": [
                "2026",
                "2026",
                "2026",
                "2026",
                "2026",
                "2026",
                "2026 ИТОГО",
                "2027 ИТОГО",
            ],
            "Квартал": [1, 1, 1, 2, 3, 4, "ИТОГ", "ИТОГ"],
            "Месяц": [1, 2, 3, 4, 7, 10, "", ""],
            "№ недели": [2, 6, 10, 15, 28, 41, "", ""],
            "№ конт": ["1201", "1202", "1203", "1204", "1205", "1206", "", ""],
            "Статус": [
                "✅ Факт",
                "⚠️ Плановая дата прошла, факта нет",
                "🟡 План",
                "🟡 План",
                "🟡 План",
                "🟡 План",
                "⚠️ Плановая дата прошла, факта нет",
                "",
            ],
            "N18": [10, 20, 25, 30, 40, 50, 175, 200],
            "N28": [1.5, 2.5, 2, 3.5, 4.5, 5.5, 19.5, 25.5],
        })
        self.workbook = load_workbook(BytesIO(export_finance_table(self.df, unit="шт")))
        self.worksheet = self.workbook["План поставок"]

    @staticmethod
    def _rgb(cell) -> str:
        return cell.fill.fgColor.rgb[-6:]

    def test_workbook_structure_and_values_are_preserved(self):
        self.assertEqual(
            [cell.value for cell in self.worksheet[1]],
            self.df.columns.tolist(),
        )
        self.assertEqual(self.worksheet.cell(2, 1).value, "2026")
        self.assertEqual(self.worksheet.cell(2, 7).value, 10)
        self.assertEqual(self.worksheet.cell(8, 1).value, "2026 ИТОГО")
        self.assertEqual(self.worksheet.cell(8, 7).value, 175)
        self.assertEqual(self.worksheet.cell(9, 1).value, "2027 ИТОГО")
        self.assertEqual(self.worksheet.cell(9, 7).value, 200)
        self.assertIsInstance(self.worksheet.cell(2, 7).value, int)
        self.assertIsInstance(self.worksheet.cell(2, 8).value, float)

    def test_header_freeze_filter_borders_and_widths(self):
        self.assertEqual(self._rgb(self.worksheet.cell(1, 1)), HEADER_FILL)
        self.assertTrue(self.worksheet.cell(1, 1).font.bold)
        self.assertEqual(self.worksheet.freeze_panes, "A2")
        self.assertEqual(self.worksheet.auto_filter.ref, "A1:H9")
        self.assertEqual(
            self.worksheet.cell(2, 1).border.left.color.rgb[-6:],
            BORDER_COLOR,
        )
        self.assertLessEqual(self.worksheet.column_dimensions["G"].width, 18)

    def test_total_fact_overdue_and_quarter_styles_match_mapping(self):
        expected_rows = {
            2: "fact",
            3: "overdue",
            4: "quarter_1",
            5: "quarter_2",
            6: "quarter_3",
            7: "quarter_4",
            8: "total",
            9: "total",
        }

        for row_number, semantic_style in expected_rows.items():
            expected = FINANCE_ROW_STYLES[semantic_style]
            cell = self.worksheet.cell(row_number, 1)
            self.assertEqual(self._rgb(cell), expected["fill"])
            self.assertEqual(cell.font.color.rgb[-6:], expected["font_color"])
            self.assertEqual(cell.font.bold, expected["bold"])

    def test_total_style_has_priority_over_status_and_quarter(self):
        total_row = self.df.iloc[6]
        self.assertEqual(
            get_finance_row_style(total_row),
            FINANCE_ROW_STYLES["total"],
        )

    def test_number_formats_follow_selected_unit_without_changing_values(self):
        expectations = {
            "шт": "#,##0",
            "кг": "#,##0.00",
            "уп": "#,##0.00",
        }

        for unit, expected_format in expectations.items():
            with self.subTest(unit=unit):
                workbook = load_workbook(
                    BytesIO(export_finance_table(self.df, unit=unit))
                )
                worksheet = workbook["План поставок"]
                self.assertEqual(worksheet.cell(2, 7).value, 10)
                self.assertEqual(worksheet.cell(2, 8).value, 1.5)
                self.assertEqual(worksheet.cell(2, 7).number_format, expected_format)
                self.assertEqual(worksheet.cell(2, 8).number_format, expected_format)


if __name__ == "__main__":
    unittest.main()
