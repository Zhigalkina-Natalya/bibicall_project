import unittest
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.osg.exporter import export_osg_table


class OsgExcelExporterTests(unittest.TestCase):
    def test_formatted_workbook_contract(self):
        df = pd.DataFrame({
            "Категория": ["ПЮРЕ МЯСНЫЕ консервы"],
            "SKU": ["Пюре 80 г. Кролик в сливочном соусе"],
            "Срок годности": ["27.08.2027"],
            "ОСГ %": [50.958904],
            "Риск": ["Пограничный 🟠"],
            "Остаток": [281.0],
            "Зарезервировано": [252.0],
            "Свободный остаток": [29.0],
            "Доля резерва, %": [89.679715],
        }, index=[48])

        content = export_osg_table(df, "Срочные продажи")
        workbook = load_workbook(BytesIO(content))
        worksheet = workbook["Срочные продажи"]

        self.assertEqual(worksheet.max_row, 2)
        self.assertEqual(worksheet.max_column, len(df.columns))
        self.assertEqual(
            [cell.value for cell in worksheet[1]],
            df.columns.tolist(),
        )
        self.assertNotIn("index", [cell.value for cell in worksheet[1]])
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(worksheet.auto_filter.ref, worksheet.dimensions)

        columns = {cell.value: cell.column for cell in worksheet[1]}
        self.assertEqual(
            worksheet.cell(2, columns["Срок годности"]).number_format,
            "DD.MM.YYYY",
        )
        self.assertEqual(
            worksheet.cell(2, columns["Остаток"]).number_format,
            "#,##0",
        )
        self.assertEqual(
            worksheet.cell(2, columns["Доля резерва, %"]).number_format,
            '0.##"%"',
        )
        self.assertEqual(
            worksheet.cell(2, columns["ОСГ %"]).fill.fill_type,
            "solid",
        )


if __name__ == "__main__":
    unittest.main()
